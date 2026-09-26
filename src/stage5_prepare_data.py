"""Stage 5 confirmatory data acquisition and validation.

This script performs DATA PREPARATION ONLY.

It does NOT:
- run coordinate descent;
- run ASCD;
- inspect Stage 5 optimization outcomes;
- alter Stage 3 or Stage 4 results.

Real datasets
-------------
1. California Housing
   sklearn.datasets.fetch_california_housing

2. Concrete Compressive Strength
   UCI dataset ID 165 through ucimlrepo

3. Airfoil Self-Noise
   UCI dataset ID 291 through ucimlrepo

4. Residential Building
   UCI dataset ID 437 loaded directly from the official UCI ZIP archive.

The direct Residential loader is necessary because UCI dataset 437 exists
but is currently not importable through ucimlrepo.

Residential Building workbook audit
-----------------------------------
The official workbook contains:

    sheets:
        Data
        Descriptions

The Data sheet contains:
    372 observations
    109 columns
    four project-date variables
    V-1 through V-8
    V-11 through V-29 repeated across time lags 1 through 5
    V-9 and V-10 as outputs

The Descriptions sheet identifies:
    V-9  = Actual sales prices (output)
    V-10 = Actual construction costs (output)

Stage 5 therefore:
    target = V-9
    excludes V-9 and V-10 from X

This leaves 107 predictor columns.

Real-data preprocessing
-----------------------
- fixed 80/20 train/test split;
- split seed = 20260926;
- training feature mean/std only;
- zero-variance training features removed and recorded;
- response centered/scaled using training statistics only;
- test data transformed with training statistics only.

Held-out synthetic data
-----------------------
- n = 800;
- d = 96;
- Toeplitz covariance Sigma_ij = rho ** |i-j|;
- rho in {0.60, 0.90};
- balanced or geometric feature scaling;
- problem seeds 100..111;
- beta ~ N(0, I), normalized to unit Euclidean norm;
- Gaussian observation noise sigma = 0.15.

Outputs
-------
results/stage5_confirmatory/
    dataset_manifest.json
    prepared_data/
        california_housing.npz
        concrete_compressive_strength.npz
        airfoil_self_noise.npz
        residential_building.npz
        heldout_balanced_moderate_seed100.npz
        ...
        heldout_scaled_strong_seed111.npz

Local cache
-----------
data/stage5_cache/
    raw/
    downloads/
    sklearn/

Run
---
From repository root:

    python src/stage5_prepare_data.py

To force a re-download of external datasets:

    python src/stage5_prepare_data.py --refresh
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import platform
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from urllib.request import Request, urlopen
from zipfile import ZipFile

import numpy as np
import pandas as pd
import sklearn
from sklearn.datasets import fetch_california_housing
from sklearn.model_selection import train_test_split


# =============================================================================
# Repository paths
# =============================================================================

ROOT = Path(__file__).resolve().parents[1]

RESULT_ROOT = (
    ROOT
    / "results"
    / "stage5_confirmatory"
)

PREPARED_ROOT = (
    RESULT_ROOT
    / "prepared_data"
)

CACHE_ROOT = (
    ROOT
    / "data"
    / "stage5_cache"
)

RAW_CACHE_ROOT = (
    CACHE_ROOT
    / "raw"
)

DOWNLOAD_ROOT = (
    CACHE_ROOT
    / "downloads"
)

SKLEARN_CACHE_ROOT = (
    CACHE_ROOT
    / "sklearn"
)

MANIFEST_PATH = (
    RESULT_ROOT
    / "dataset_manifest.json"
)


# =============================================================================
# Frozen Stage 5 protocol constants
# =============================================================================

SPLIT_SEED = 20260926

TRAIN_FRACTION = 0.80
TEST_FRACTION = 0.20


# -----------------------------------------------------------------------------
# Held-out synthetic protocol
# -----------------------------------------------------------------------------

SYNTHETIC_N = 800
SYNTHETIC_D = 96

SYNTHETIC_SEEDS = tuple(
    range(100, 112)
)

SYNTHETIC_NOISE_STD = 0.15

SYNTHETIC_BETA_DISTRIBUTION = (
    "standard_normal_then_unit_l2_normalized"
)

SYNTHETIC_SCALE_MIN_EXP = -1.5
SYNTHETIC_SCALE_MAX_EXP = 1.5

SYNTHETIC_CASES = (
    {
        "name": "heldout_balanced_moderate",
        "rho": 0.60,
        "scaled": False,
    },
    {
        "name": "heldout_balanced_strong",
        "rho": 0.90,
        "scaled": False,
    },
    {
        "name": "heldout_scaled_moderate",
        "rho": 0.60,
        "scaled": True,
    },
    {
        "name": "heldout_scaled_strong",
        "rho": 0.90,
        "scaled": True,
    },
)


# =============================================================================
# External dataset declarations
# =============================================================================

RESIDENTIAL_ARCHIVE_URL = (
    "https://archive.ics.uci.edu/static/public/437/"
    "residential%2Bbuilding%2Bdata%2Bset.zip"
)

RESIDENTIAL_ARCHIVE_NAME = (
    "residential_building_uci_437.zip"
)

REAL_DATASETS = {
    "california_housing": {
        "display_name": (
            "California Housing"
        ),
        "source": "scikit-learn",
        "source_id": (
            "fetch_california_housing"
        ),
        "source_url": (
            "https://scikit-learn.org/stable/modules/"
            "generated/sklearn.datasets."
            "fetch_california_housing.html"
        ),
    },
    "concrete_compressive_strength": {
        "display_name": (
            "Concrete Compressive Strength"
        ),
        "source": (
            "UCI Machine Learning Repository"
        ),
        "source_id": 165,
        "doi": "10.24432/C5PK67",
        "source_url": (
            "https://archive.ics.uci.edu/"
            "dataset/165/"
            "concrete+compressive+strength"
        ),
    },
    "airfoil_self_noise": {
        "display_name": (
            "Airfoil Self-Noise"
        ),
        "source": (
            "UCI Machine Learning Repository"
        ),
        "source_id": 291,
        "doi": "10.24432/C5VW2C",
        "source_url": (
            "https://archive.ics.uci.edu/"
            "dataset/291/"
            "airfoil+self+noise"
        ),
    },
    "residential_building": {
        "display_name": (
            "Residential Building"
        ),
        "source": (
            "UCI Machine Learning Repository"
        ),
        "source_id": 437,
        "doi": "10.24432/C5S896",
        "source_url": (
            "https://archive.ics.uci.edu/"
            "dataset/437/"
            "residential+building+data+set"
        ),
        "archive_url": (
            RESIDENTIAL_ARCHIVE_URL
        ),
    },
}


# =============================================================================
# Expected source schemas
# =============================================================================

EXPECTED_REAL_SCHEMAS = {
    "california_housing": {
        "rows": 20640,
        "features": 8,
        "targets": 1,
    },
    "concrete_compressive_strength": {
        "rows": 1030,
        "features": 8,
        "targets": 1,
    },
    "airfoil_self_noise": {
        "rows": 1503,
        "features": 5,
        "targets": 1,
    },
    "residential_building": {
        # Official workbook:
        # 109 columns =
        #   4 project-date columns
        #   + V-1 ... V-8
        #   + V-11 ... V-29 repeated for time lags 1 ... 5
        #   + V-9 and V-10 outputs
        #
        # V-9  = Actual sales prices (output)
        # V-10 = Actual construction costs (output)
        #
        # 109 - 2 outputs = 107 predictors.
        "rows": 372,
        "features": 107,
        "targets": 2,
        "workbook_columns": 109,
    },
}


# =============================================================================
# Data container
# =============================================================================

@dataclass
class RawRealDataset:
    slug: str

    X: np.ndarray
    Y: np.ndarray

    feature_names: list[str]
    target_names: list[str]

    target_descriptions: dict[str, str]

    source_metadata: dict[str, Any]

    loaded_from_cache: bool


# =============================================================================
# Utility functions
# =============================================================================

def ensure_directories() -> None:
    """Create all Stage 5 output/cache directories."""
    for path in (
        RESULT_ROOT,
        PREPARED_ROOT,
        CACHE_ROOT,
        RAW_CACHE_ROOT,
        DOWNLOAD_ROOT,
        SKLEARN_CACHE_ROOT,
    ):
        path.mkdir(
            parents=True,
            exist_ok=True,
        )


def utc_now_string() -> str:
    return (
        datetime
        .now(timezone.utc)
        .replace(microsecond=0)
        .isoformat()
    )


def package_version(
    package_name: str,
) -> str | None:
    try:
        return importlib.metadata.version(
            package_name
        )
    except importlib.metadata.PackageNotFoundError:
        return None


def json_safe(
    value: Any,
) -> Any:
    """Convert common metadata values to JSON-safe forms."""
    if value is None:
        return None

    if isinstance(
        value,
        (
            str,
            int,
            float,
            bool,
        ),
    ):
        return value

    if isinstance(
        value,
        np.generic,
    ):
        return value.item()

    if isinstance(
        value,
        Path,
    ):
        return str(value)

    if isinstance(
        value,
        dict,
    ):
        return {
            str(key): json_safe(item)
            for key, item
            in value.items()
        }

    if isinstance(
        value,
        (
            list,
            tuple,
        ),
    ):
        return [
            json_safe(item)
            for item in value
        ]

    return str(value)


def sha256_bytes(
    payload: bytes,
) -> str:
    return hashlib.sha256(
        payload
    ).hexdigest()


def sha256_file(
    path: Path,
) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for block in iter(
            lambda: handle.read(
                1024 * 1024
            ),
            b"",
        ):
            digest.update(block)

    return digest.hexdigest()


def _hash_text(
    digest: Any,
    value: str,
) -> None:
    encoded = value.encode(
        "utf-8"
    )

    digest.update(
        len(encoded).to_bytes(
            8,
            "little",
        )
    )

    digest.update(encoded)


def hash_arrays_and_labels(
    arrays: Iterable[np.ndarray],
    labels: Iterable[str] = (),
) -> str:
    """Canonical content hash independent of ZIP timestamps."""
    digest = hashlib.sha256()

    for array in arrays:
        arr = np.ascontiguousarray(
            array
        )

        _hash_text(
            digest,
            str(arr.dtype),
        )

        _hash_text(
            digest,
            json.dumps(
                arr.shape
            ),
        )

        digest.update(
            arr.tobytes(
                order="C"
            )
        )

    for label in labels:
        _hash_text(
            digest,
            str(label),
        )

    return digest.hexdigest()


def assert_finite(
    name: str,
    array: np.ndarray,
) -> None:
    if not np.issubdtype(
        array.dtype,
        np.number,
    ):
        raise TypeError(
            f"{name} is not numeric: "
            f"dtype={array.dtype}"
        )

    bad = ~np.isfinite(
        array
    )

    if np.any(bad):
        locations = (
            np.argwhere(bad)
            [:10]
            .tolist()
        )

        raise ValueError(
            f"{name} contains "
            f"{int(np.sum(bad))} "
            "non-finite values. "
            f"First locations: {locations}"
        )


def normalize_name(
    value: Any,
) -> str:
    """Normalize identifiers for robust matching."""
    return re.sub(
        r"[^a-z0-9]+",
        "",
        str(value).lower(),
    )


def normalize_variable_id(
    value: Any,
) -> str:
    """Normalize V-9, V9, V_9 etc. to v9."""
    return normalize_name(
        value
    )


def build_residential_unique_column_names(
    group_header: list[Any],
    variable_header: list[Any],
) -> list[str]:
    """Create unique, auditable Residential Building column names.

    The official workbook repeats variables V-11 through V-29
    across time lags 1 through 5.

    Naming policy
    -------------
    Project-date variables:
        START_YEAR
        START_QUARTER
        COMPLETION_YEAR
        COMPLETION_QUARTER

    Project physical/financial variables:
        V-1 through V-8

    Economic variables:
        V-11_lag1 through V-29_lag5

    Outputs:
        V-9
        V-10

    The function deliberately fails if the resulting identifiers
    are not unique.
    """

    if len(group_header) != len(variable_header):
        raise ValueError(
            "Residential header rows have different lengths: "
            f"{len(group_header)} vs {len(variable_header)}."
        )

    unique_names: list[str] = []

    current_group = ""

    for column_index, (
        group_value,
        variable_value,
    ) in enumerate(
        zip(
            group_header,
            variable_header,
        ),
        start=1,
    ):
        if group_value is not None:
            current_group = str(
                group_value
            ).strip()

        if variable_value is None:
            raise RuntimeError(
                "Residential second header row contains "
                f"a blank variable at column {column_index}."
            )

        variable = str(
            variable_value
        ).strip()

        if not variable:
            raise RuntimeError(
                "Residential second header row contains "
                f"an empty variable at column {column_index}."
            )

        group_upper = (
            current_group.upper()
        )

        lag_match = re.search(
            r"TIME\s+LAG\s+(\d+)",
            group_upper,
        )

        if lag_match is not None:
            lag = int(
                lag_match.group(1)
            )

            if lag not in {
                1,
                2,
                3,
                4,
                5,
            }:
                raise RuntimeError(
                    "Unexpected Residential time lag "
                    f"{lag} at column {column_index}."
                )

            unique_name = (
                f"{variable}_lag{lag}"
            )

        elif variable in {
            "START YEAR",
            "START QUARTER",
            "COMPLETION YEAR",
            "COMPLETION QUARTER",
        }:
            unique_name = (
                variable.replace(
                    " ",
                    "_",
                )
            )

        else:
            unique_name = variable

        unique_names.append(
            unique_name
        )

    if len(unique_names) != 109:
        raise RuntimeError(
            "Residential unique-name construction "
            f"produced {len(unique_names)} names; "
            "expected 109."
        )

    if len(set(unique_names)) != len(unique_names):
        duplicates = sorted(
            {
                name
                for name in unique_names
                if unique_names.count(name) > 1
            }
        )

        raise RuntimeError(
            "Residential unique-name construction failed. "
            f"Duplicate names remain: {duplicates}"
        )

    required_names = {
        "START_YEAR",
        "START_QUARTER",
        "COMPLETION_YEAR",
        "COMPLETION_QUARTER",
        "V-1",
        "V-8",
        "V-11_lag1",
        "V-29_lag1",
        "V-11_lag5",
        "V-29_lag5",
        "V-9",
        "V-10",
    }

    missing = (
        required_names
        - set(unique_names)
    )

    if missing:
        raise RuntimeError(
            "Residential unique-name construction "
            "is missing expected identifiers: "
            f"{sorted(missing)}"
        )

    return unique_names


# =============================================================================
# pandas / ucimlrepo helpers
# =============================================================================

def dataframe_to_numeric(
    frame: pd.DataFrame,
    dataset_name: str,
    role: str,
) -> np.ndarray:
    if frame is None:
        raise ValueError(
            f"{dataset_name}: "
            f"{role} dataframe is None."
        )

    if frame.empty:
        raise ValueError(
            f"{dataset_name}: "
            f"{role} dataframe is empty."
        )

    if frame.isnull().any().any():
        missing = (
            frame
            .isnull()
            .sum()
        )

        missing = (
            missing[
                missing > 0
            ]
            .to_dict()
        )

        raise ValueError(
            f"{dataset_name}: "
            f"missing values found in "
            f"{role}: {missing}"
        )

    try:
        numeric = frame.apply(
            pd.to_numeric,
            errors="raise",
        )
    except Exception as exc:
        raise TypeError(
            f"{dataset_name}: "
            f"non-numeric values found "
            f"in {role}."
        ) from exc

    array = numeric.to_numpy(
        dtype=np.float64,
        copy=True,
    )

    assert_finite(
        f"{dataset_name} {role}",
        array,
    )

    return array


def import_ucimlrepo():
    try:
        from ucimlrepo import (
            fetch_ucirepo,
        )
    except ImportError as exc:
        raise RuntimeError(
            "Stage 5 requires ucimlrepo "
            "for UCI datasets 165 and 291.\n"
            "\nInstall with:\n"
            "    python -m pip install "
            "ucimlrepo"
        ) from exc

    return fetch_ucirepo


def extract_uci_target_descriptions(
    variables: Any,
) -> dict[str, str]:
    descriptions: dict[
        str,
        str,
    ] = {}

    if not isinstance(
        variables,
        pd.DataFrame,
    ):
        return descriptions

    normalized_columns = {
        normalize_name(column): column
        for column
        in variables.columns
    }

    role_column = (
        normalized_columns.get(
            "role"
        )
    )

    name_column = (
        normalized_columns.get(
            "name"
        )
        or normalized_columns.get(
            "variablename"
        )
    )

    description_column = (
        normalized_columns.get(
            "description"
        )
        or normalized_columns.get(
            "variabledescription"
        )
    )

    if (
        role_column is None
        or name_column is None
    ):
        return descriptions

    for _, row in (
        variables.iterrows()
    ):
        role = normalize_name(
            row.get(
                role_column,
                "",
            )
        )

        if "target" not in role:
            continue

        raw_name = row.get(
            name_column,
            "",
        )

        if pd.isna(raw_name):
            continue

        name = str(
            raw_name
        ).strip()

        if not name:
            continue

        description = ""

        if (
            description_column
            is not None
        ):
            raw_description = row.get(
                description_column,
                "",
            )

            if pd.notna(
                raw_description
            ):
                description = str(
                    raw_description
                ).strip()

        descriptions[
            name
        ] = description

    return descriptions


# =============================================================================
# Raw cache
# =============================================================================

def raw_cache_path(
    slug: str,
) -> Path:
    return (
        RAW_CACHE_ROOT
        / f"{slug}.npz"
    )


def save_raw_cache(
    dataset: RawRealDataset,
) -> None:
    path = raw_cache_path(
        dataset.slug
    )

    description_names = list(
        dataset
        .target_descriptions
        .keys()
    )

    description_values = [
        dataset
        .target_descriptions[name]
        for name
        in description_names
    ]

    metadata_json = json.dumps(
        json_safe(
            dataset.source_metadata
        ),
        sort_keys=True,
    )

    np.savez_compressed(
        path,

        X=np.asarray(
            dataset.X,
            dtype=np.float64,
        ),

        Y=np.asarray(
            dataset.Y,
            dtype=np.float64,
        ),

        feature_names=np.asarray(
            dataset.feature_names,
            dtype=np.str_,
        ),

        target_names=np.asarray(
            dataset.target_names,
            dtype=np.str_,
        ),

        target_description_names=(
            np.asarray(
                description_names,
                dtype=np.str_,
            )
        ),

        target_description_values=(
            np.asarray(
                description_values,
                dtype=np.str_,
            )
        ),

        source_metadata_json=np.asarray(
            metadata_json,
            dtype=np.str_,
        ),
    )


def load_raw_cache(
    slug: str,
) -> RawRealDataset:
    path = raw_cache_path(
        slug
    )

    with np.load(
        path,
        allow_pickle=False,
    ) as payload:

        X = np.asarray(
            payload["X"],
            dtype=np.float64,
        )

        Y = np.asarray(
            payload["Y"],
            dtype=np.float64,
        )

        feature_names = [
            str(item)
            for item
            in payload[
                "feature_names"
            ].tolist()
        ]

        target_names = [
            str(item)
            for item
            in payload[
                "target_names"
            ].tolist()
        ]

        description_names = [
            str(item)
            for item
            in payload[
                "target_description_names"
            ].tolist()
        ]

        description_values = [
            str(item)
            for item
            in payload[
                "target_description_values"
            ].tolist()
        ]

        if (
            "source_metadata_json"
            in payload.files
        ):
            metadata_text = str(
                payload[
                    "source_metadata_json"
                ].item()
            )

            try:
                source_metadata = (
                    json.loads(
                        metadata_text
                    )
                )
            except json.JSONDecodeError:
                source_metadata = {
                    "cached_metadata_parse_failed":
                    True,
                }
        else:
            # Backwards compatibility with
            # raw caches generated by the
            # earlier Stage 5 script.
            source_metadata = {
                "legacy_raw_cache":
                True,
            }

    assert_finite(
        f"{slug} cached X",
        X,
    )

    assert_finite(
        f"{slug} cached Y",
        Y,
    )

    source_metadata[
        "loaded_from_local_raw_cache"
    ] = True

    return RawRealDataset(
        slug=slug,
        X=X,
        Y=Y,
        feature_names=feature_names,
        target_names=target_names,
        target_descriptions=dict(
            zip(
                description_names,
                description_values,
            )
        ),
        source_metadata=source_metadata,
        loaded_from_cache=True,
    )


# =============================================================================
# California Housing loader
# =============================================================================

def load_california(
    refresh: bool,
) -> RawRealDataset:
    slug = (
        "california_housing"
    )

    cache = raw_cache_path(
        slug
    )

    if (
        cache.is_file()
        and not refresh
    ):
        print(
            "  using local raw cache"
        )

        return load_raw_cache(
            slug
        )

    bunch = (
        fetch_california_housing(
            data_home=str(
                SKLEARN_CACHE_ROOT
            ),
            download_if_missing=True,
            as_frame=True,
        )
    )

    X_frame = bunch.data

    if isinstance(
        bunch.target,
        pd.Series,
    ):
        Y_frame = (
            bunch
            .target
            .to_frame()
        )
    else:
        Y_frame = pd.DataFrame(
            {
                "MedHouseVal":
                bunch.target,
            }
        )

    X = dataframe_to_numeric(
        X_frame,
        slug,
        "features",
    )

    Y = dataframe_to_numeric(
        Y_frame,
        slug,
        "targets",
    )

    dataset = RawRealDataset(
        slug=slug,
        X=X,
        Y=Y,
        feature_names=[
            str(column)
            for column
            in X_frame.columns
        ],
        target_names=[
            str(column)
            for column
            in Y_frame.columns
        ],
        target_descriptions={},
        source_metadata={
            "loader": (
                "sklearn.datasets."
                "fetch_california_housing"
            ),
        },
        loaded_from_cache=False,
    )

    save_raw_cache(
        dataset
    )

    return dataset


# =============================================================================
# Generic UCI loader for IDs 165 and 291
# =============================================================================

def fetch_uci_raw(
    slug: str,
    uci_id: int,
) -> RawRealDataset:
    fetch_ucirepo = (
        import_ucimlrepo()
    )

    print(
        f"  downloading UCI dataset "
        f"{uci_id} for {slug}"
    )

    dataset = (
        fetch_ucirepo(
            id=uci_id
        )
    )

    X_frame = (
        dataset.data.features
    )

    Y_frame = (
        dataset.data.targets
    )

    if isinstance(
        Y_frame,
        pd.Series,
    ):
        Y_frame = (
            Y_frame.to_frame()
        )

    if not isinstance(
        X_frame,
        pd.DataFrame,
    ):
        X_frame = pd.DataFrame(
            X_frame
        )

    if not isinstance(
        Y_frame,
        pd.DataFrame,
    ):
        Y_frame = pd.DataFrame(
            Y_frame
        )

    X = dataframe_to_numeric(
        X_frame,
        slug,
        "features",
    )

    Y = dataframe_to_numeric(
        Y_frame,
        slug,
        "targets",
    )

    target_descriptions = (
        extract_uci_target_descriptions(
            getattr(
                dataset,
                "variables",
                None,
            )
        )
    )

    result = RawRealDataset(
        slug=slug,
        X=X,
        Y=Y,
        feature_names=[
            str(column)
            for column
            in X_frame.columns
        ],
        target_names=[
            str(column)
            for column
            in Y_frame.columns
        ],
        target_descriptions=(
            target_descriptions
        ),
        source_metadata={
            "uci_id": uci_id,
            "returned_metadata": (
                json_safe(
                    getattr(
                        dataset,
                        "metadata",
                        None,
                    )
                )
            ),
        },
        loaded_from_cache=False,
    )

    return result


def load_uci_dataset(
    slug: str,
    uci_id: int,
    refresh: bool,
) -> RawRealDataset:
    cache = raw_cache_path(
        slug
    )

    if (
        cache.is_file()
        and not refresh
    ):
        print(
            "  using local raw cache"
        )

        return load_raw_cache(
            slug
        )

    dataset = fetch_uci_raw(
        slug=slug,
        uci_id=uci_id,
    )

    save_raw_cache(
        dataset
    )

    return dataset


# =============================================================================
# Residential Building direct UCI loader
# =============================================================================

def import_openpyxl():
    try:
        from openpyxl import (
            load_workbook,
        )
    except ImportError as exc:
        raise RuntimeError(
            "Residential Building requires "
            "openpyxl.\n\n"
            "Install with:\n"
            "    python -m pip install "
            "openpyxl"
        ) from exc

    return load_workbook


def download_residential_archive(
    refresh: bool,
) -> tuple[bytes, Path]:
    archive_path = (
        DOWNLOAD_ROOT
        / RESIDENTIAL_ARCHIVE_NAME
    )

    if (
        archive_path.is_file()
        and not refresh
    ):
        print(
            "  using cached official UCI ZIP"
        )

        return (
            archive_path.read_bytes(),
            archive_path,
        )

    print(
        "  downloading official "
        "UCI Residential Building archive"
    )

    request = Request(
        RESIDENTIAL_ARCHIVE_URL,
        headers={
            "User-Agent": (
                "Mozilla/5.0 "
                "Stage5-Research-Preflight/1.0"
            ),
        },
    )

    try:
        with urlopen(
            request,
            timeout=60,
        ) as response:
            payload = response.read()
    except Exception as exc:
        raise RuntimeError(
            "Failed to download official "
            "Residential Building UCI archive."
        ) from exc

    if not payload:
        raise RuntimeError(
            "Downloaded Residential Building "
            "archive is empty."
        )

    archive_path.write_bytes(
        payload
    )

    return (
        payload,
        archive_path,
    )


def parse_residential_descriptions(
    worksheet: Any,
) -> dict[str, str]:
    """Parse V-* descriptions from the workbook."""
    descriptions: dict[
        str,
        str,
    ] = {}

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if len(row) < 3:
            continue

        raw_variable = row[1]
        raw_description = row[2]

        if raw_variable is None:
            continue

        variable = str(
            raw_variable
        ).strip()

        if not variable:
            continue

        description = (
            ""
            if raw_description is None
            else str(
                raw_description
            ).strip()
        )

        descriptions[
            variable
        ] = description

    return descriptions


def is_sale_price_description(
    text: str,
) -> bool:
    normalized = normalize_name(
        text
    )

    return (
        (
            "sale" in normalized
            or "sales" in normalized
        )
        and "price" in normalized
        and "output" in normalized
    )


def is_construction_cost_description(
    text: str,
) -> bool:
    normalized = normalize_name(
        text
    )

    return (
        "construction" in normalized
        and "cost" in normalized
        and "output" in normalized
    )


def load_residential_direct(
    refresh: bool,
) -> RawRealDataset:
    """Load UCI 437 directly from its official workbook."""
    slug = (
        "residential_building"
    )

    cache = raw_cache_path(
        slug
    )

    if (
        cache.is_file()
        and not refresh
    ):
        print(
            "  using local raw cache"
        )

        cached = load_raw_cache(
            slug
        )

        # Do not silently accept a stale cache made
        # under an earlier interpretation.
        expected = (
            EXPECTED_REAL_SCHEMAS[
                slug
            ]
        )

        if (
            cached.X.shape
            != (
                expected["rows"],
                expected["features"],
            )
        ):
            raise RuntimeError(
                "Existing Residential Building "
                "raw cache has incompatible shape "
                f"{cached.X.shape}. Expected "
                f"({expected['rows']}, "
                f"{expected['features']}).\n\n"
                "Delete the stale cache or rerun:\n"
                "    python "
                "src/stage5_prepare_data.py "
                "--refresh"
            )

        if (
            cached.Y.shape
            != (
                expected["rows"],
                expected["targets"],
            )
        ):
            raise RuntimeError(
                "Existing Residential Building "
                "target cache has incompatible "
                f"shape {cached.Y.shape}."
            )

        if len(
            cached.feature_names
        ) != expected["features"]:
            raise RuntimeError(
                "Existing Residential raw cache has "
                f"{len(cached.feature_names)} feature names; "
                f"expected {expected['features']}.\n\n"
                "Rerun with:\n"
                "    python "
                "src/stage5_prepare_data.py "
                "--refresh"
            )

        if len(
            set(cached.feature_names)
        ) != len(
            cached.feature_names
        ):
            raise RuntimeError(
                "Existing Residential raw cache contains "
                "non-unique predictor names from the earlier "
                "header interpretation.\n\n"
                "Rerun with:\n"
                "    python "
                "src/stage5_prepare_data.py "
                "--refresh"
            )

        required_feature_names = {
            "V-11_lag1",
            "V-29_lag1",
            "V-11_lag5",
            "V-29_lag5",
        }

        if not required_feature_names.issubset(
            set(cached.feature_names)
        ):
            raise RuntimeError(
                "Residential cached predictor labels do not "
                "match the audited lag-aware schema.\n\n"
                "Rerun with:\n"
                "    python "
                "src/stage5_prepare_data.py "
                "--refresh"
            )

        if (
            "V-9" in cached.feature_names
            or "V-10" in cached.feature_names
        ):
            raise RuntimeError(
                "Residential cached predictors contain "
                "one or both output variables V-9/V-10. "
                "This violates the anti-leakage rule."
            )

        return cached

    load_workbook = (
        import_openpyxl()
    )

    (
        archive_bytes,
        archive_path,
    ) = download_residential_archive(
        refresh=refresh
    )

    archive_sha256 = (
        sha256_bytes(
            archive_bytes
        )
    )

    with ZipFile(
        BytesIO(
            archive_bytes
        )
    ) as zip_file:

        archive_members = (
            zip_file.namelist()
        )

        xlsx_members = [
            member
            for member
            in archive_members
            if member
            .lower()
            .endswith(".xlsx")
        ]

        if len(
            xlsx_members
        ) != 1:
            raise RuntimeError(
                "Expected exactly one XLSX "
                "file in Residential archive; "
                f"found {xlsx_members}"
            )

        workbook_member = (
            xlsx_members[0]
        )

        workbook_bytes = (
            zip_file.read(
                workbook_member
            )
        )

    workbook_sha256 = (
        sha256_bytes(
            workbook_bytes
        )
    )

    workbook = load_workbook(
        BytesIO(
            workbook_bytes
        ),
        read_only=True,
        data_only=True,
    )

    required_sheets = {
        "Data",
        "Descriptions",
    }

    if not required_sheets.issubset(
        set(
            workbook.sheetnames
        )
    ):
        raise RuntimeError(
            "Residential workbook schema "
            "changed. Expected sheets "
            f"{sorted(required_sheets)}, "
            f"found {workbook.sheetnames}."
        )

    data_sheet = (
        workbook["Data"]
    )

    descriptions_sheet = (
        workbook[
            "Descriptions"
        ]
    )

    # The inspected official workbook has:
    # rows 1-2 = two header rows
    # rows 3-374 = 372 observations
    # 109 columns total.
    if data_sheet.max_row != 374:
        raise RuntimeError(
            "Residential Data sheet row "
            "count changed: expected 374 "
            f"including two headers; found "
            f"{data_sheet.max_row}."
        )

    if data_sheet.max_column != 109:
        raise RuntimeError(
            "Residential Data sheet column "
            "count changed: expected 109; "
            f"found {data_sheet.max_column}."
        )

    group_rows = list(
        data_sheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True,
        )
    )

    variable_rows = list(
        data_sheet.iter_rows(
            min_row=2,
            max_row=2,
            values_only=True,
        )
    )

    if len(group_rows) != 1:
        raise RuntimeError(
            "Could not read Residential "
            "first header row."
        )

    if len(variable_rows) != 1:
        raise RuntimeError(
            "Could not read Residential "
            "second header row."
        )

    group_header = list(
        group_rows[0]
    )

    variable_header = list(
        variable_rows[0]
    )

    header = (
        build_residential_unique_column_names(
            group_header,
            variable_header,
        )
    )

    if len(header) != 109:
        raise RuntimeError(
            "Residential second header row "
            f"has {len(header)} columns, "
            "expected 109."
        )

    blank_columns = [
        index + 1
        for index, value
        in enumerate(header)
        if (
            value is None
            or value == ""
        )
    ]

    if blank_columns:
        raise RuntimeError(
            "Residential second header row "
            "contains blank column names at "
            f"positions {blank_columns}."
        )

    normalized_header = {
        normalize_variable_id(name):
        index
        for index, name
        in enumerate(header)
    }

    if "v9" not in normalized_header:
        raise RuntimeError(
            "Residential workbook does not "
            "contain expected V-9 column."
        )

    if "v10" not in normalized_header:
        raise RuntimeError(
            "Residential workbook does not "
            "contain expected V-10 column."
        )

    sale_index = (
        normalized_header["v9"]
    )

    construction_index = (
        normalized_header["v10"]
    )

    descriptions = (
        parse_residential_descriptions(
            descriptions_sheet
        )
    )

    normalized_descriptions = {
        normalize_variable_id(key):
        value
        for key, value
        in descriptions.items()
    }

    sale_description = (
        normalized_descriptions.get(
            "v9",
            "",
        )
    )

    construction_description = (
        normalized_descriptions.get(
            "v10",
            "",
        )
    )

    if not is_sale_price_description(
        sale_description
    ):
        raise RuntimeError(
            "Residential V-9 description "
            "does not unambiguously identify "
            "a sales-price output.\n"
            f"Observed description: "
            f"{sale_description!r}"
        )

    if not is_construction_cost_description(
        construction_description
    ):
        raise RuntimeError(
            "Residential V-10 description "
            "does not unambiguously identify "
            "a construction-cost output.\n"
            f"Observed description: "
            f"{construction_description!r}"
        )

    raw_rows = list(
        data_sheet.iter_rows(
            min_row=3,
            max_row=374,
            values_only=True,
        )
    )

    if len(raw_rows) != 372:
        raise RuntimeError(
            "Residential workbook yielded "
            f"{len(raw_rows)} observations, "
            "expected 372."
        )

    try:
        full_matrix = np.asarray(
            raw_rows,
            dtype=np.float64,
        )
    except (
        TypeError,
        ValueError,
    ) as exc:
        raise TypeError(
            "Residential workbook contains "
            "a non-numeric value in its data "
            "rows."
        ) from exc

    if full_matrix.shape != (
        372,
        109,
    ):
        raise RuntimeError(
            "Residential numeric matrix "
            "has unexpected shape "
            f"{full_matrix.shape}; "
            "expected (372, 109)."
        )

    assert_finite(
        "Residential full workbook matrix",
        full_matrix,
    )

    # -----------------------------------------------------------------
    # Critical anti-leakage rule:
    #
    # V-9  = actual sales price target
    # V-10 = actual construction cost output
    #
    # BOTH outputs are excluded from predictors.
    # -----------------------------------------------------------------

    output_indices = {
        sale_index,
        construction_index,
    }

    feature_indices = [
        index
        for index
        in range(
            full_matrix.shape[1]
        )
        if index
        not in output_indices
    ]

    X = full_matrix[
        :,
        feature_indices,
    ]

    Y = full_matrix[
        :,
        [
            sale_index,
            construction_index,
        ],
    ]

    feature_names = [
        header[index]
        for index
        in feature_indices
    ]

    target_names = [
        header[sale_index],
        header[
            construction_index
        ],
    ]

    target_descriptions = {
        header[sale_index]:
        sale_description,

        header[construction_index]:
        construction_description,
    }

    expected = (
        EXPECTED_REAL_SCHEMAS[
            slug
        ]
    )

    if X.shape != (
        expected["rows"],
        expected["features"],
    ):
        raise RuntimeError(
            "Residential predictor matrix "
            f"has shape {X.shape}; expected "
            f"({expected['rows']}, "
            f"{expected['features']})."
        )

    if Y.shape != (
        expected["rows"],
        expected["targets"],
    ):
        raise RuntimeError(
            "Residential output matrix "
            f"has shape {Y.shape}; expected "
            f"({expected['rows']}, "
            f"{expected['targets']})."
        )

    dataset = RawRealDataset(
        slug=slug,
        X=X,
        Y=Y,
        feature_names=(
            feature_names
        ),
        target_names=(
            target_names
        ),
        target_descriptions=(
            target_descriptions
        ),
        source_metadata={
            "uci_id": 437,
            "loading_method": (
                "direct_official_uci_zip"
            ),
            "archive_url": (
                RESIDENTIAL_ARCHIVE_URL
            ),
            "archive_file": str(
                archive_path
                .relative_to(ROOT)
            ),
            "archive_sha256": (
                archive_sha256
            ),
            "xlsx_member": (
                workbook_member
            ),
            "xlsx_sha256": (
                workbook_sha256
            ),
            "workbook_sheets": (
                workbook.sheetnames
            ),
            "data_sheet_rows_including_headers":
            data_sheet.max_row,
            "data_sheet_columns": (
                data_sheet.max_column
            ),
            "data_rows": 372,
            "full_column_count": 109,
            "predictor_count_after_excluding_outputs":
            107,
            "feature_naming_rule": (
                "project-date variables use underscore names; "
                "V-1 through V-8 retain source IDs; "
                "repeated V-11 through V-29 variables are named "
                "V-<id>_lag<1..5>"
            ),
            "lagged_predictor_groups": {
                "lag1": "V-11 through V-29",
                "lag2": "V-11 through V-29",
                "lag3": "V-11 through V-29",
                "lag4": "V-11 through V-29",
                "lag5": "V-11 through V-29",
            },
            "sale_price_variable": (
                header[sale_index]
            ),
            "sale_price_description": (
                sale_description
            ),
            "construction_cost_variable": (
                header[
                    construction_index
                ]
            ),
            "construction_cost_description": (
                construction_description
            ),
            "excluded_output_variables": [
                header[sale_index],
                header[
                    construction_index
                ],
            ],
        },
        loaded_from_cache=False,
    )

    save_raw_cache(
        dataset
    )

    return dataset


# =============================================================================
# Target selection
# =============================================================================

def resolve_single_target(
    dataset: RawRealDataset,
) -> tuple[
    np.ndarray,
    str,
    dict[str, Any],
]:
    if dataset.Y.ndim != 2:
        raise ValueError(
            f"{dataset.slug}: expected "
            "2-D target matrix, got "
            f"{dataset.Y.shape}."
        )

    if dataset.Y.shape[1] != 1:
        raise ValueError(
            f"{dataset.slug}: expected "
            "exactly one target, got "
            f"{dataset.Y.shape[1]}: "
            f"{dataset.target_names}"
        )

    return (
        dataset.Y[
            :,
            0,
        ].copy(),
        dataset.target_names[0],
        {
            "resolution": (
                "single_returned_target"
            ),
        },
    )


def resolve_residential_sale_target(
    dataset: RawRealDataset,
) -> tuple[
    np.ndarray,
    str,
    dict[str, Any],
]:
    """Select V-9 only after validating its source description."""
    if dataset.Y.shape != (
        372,
        2,
    ):
        raise ValueError(
            "Residential output matrix "
            "must have shape (372, 2); "
            f"found {dataset.Y.shape}."
        )

    candidate_indices: list[
        int
    ] = []

    for index, name in enumerate(
        dataset.target_names
    ):
        description = (
            dataset
            .target_descriptions
            .get(
                name,
                "",
            )
        )

        if is_sale_price_description(
            description
        ):
            candidate_indices.append(
                index
            )

    if len(
        candidate_indices
    ) != 1:
        raise RuntimeError(
            "Residential sale-price target "
            "could not be identified "
            "unambiguously.\n"
            f"Targets: "
            f"{dataset.target_names}\n"
            f"Descriptions: "
            f"{dataset.target_descriptions}"
        )

    index = (
        candidate_indices[0]
    )

    selected_name = (
        dataset.target_names[
            index
        ]
    )

    # Additional defensive check.
    if normalize_variable_id(
        selected_name
    ) != "v9":
        raise RuntimeError(
            "Sales-price description "
            "matched an unexpected variable "
            f"{selected_name!r}; expected V-9."
        )

    return (
        dataset.Y[
            :,
            index,
        ].copy(),
        selected_name,
        {
            "resolution": (
                "official_workbook_description"
            ),
            "selected_variable": (
                selected_name
            ),
            "description": (
                dataset
                .target_descriptions[
                    selected_name
                ]
            ),
            "excluded_other_output":
            "V-10",
        },
    )


# =============================================================================
# Real-data preprocessing
# =============================================================================

def prepare_real_dataset(
    dataset: RawRealDataset,
    y_raw: np.ndarray,
    target_name: str,
    target_resolution: dict[str, Any],
) -> dict[str, Any]:

    X_raw = np.asarray(
        dataset.X,
        dtype=np.float64,
    )

    y_raw = np.asarray(
        y_raw,
        dtype=np.float64,
    ).reshape(-1)

    if len(
        dataset.feature_names
    ) != X_raw.shape[1]:
        raise ValueError(
            f"{dataset.slug}: feature-name count "
            f"{len(dataset.feature_names)} does not match "
            f"X column count {X_raw.shape[1]}."
        )

    if len(
        set(dataset.feature_names)
    ) != len(
        dataset.feature_names
    ):
        duplicates = sorted(
            {
                name
                for name in dataset.feature_names
                if dataset.feature_names.count(name) > 1
            }
        )

        raise ValueError(
            f"{dataset.slug}: predictor names are not unique: "
            f"{duplicates}"
        )

    if (
        X_raw.shape[0]
        != y_raw.shape[0]
    ):
        raise ValueError(
            f"{dataset.slug}: X/y "
            "row mismatch: "
            f"{X_raw.shape[0]} vs "
            f"{y_raw.shape[0]}."
        )

    assert_finite(
        f"{dataset.slug} X raw",
        X_raw,
    )

    assert_finite(
        f"{dataset.slug} y raw",
        y_raw,
    )

    n = X_raw.shape[0]

    indices = np.arange(
        n,
        dtype=np.int64,
    )

    (
        train_idx,
        test_idx,
    ) = train_test_split(
        indices,
        test_size=TEST_FRACTION,
        random_state=SPLIT_SEED,
        shuffle=True,
    )

    train_idx = np.asarray(
        train_idx,
        dtype=np.int64,
    )

    test_idx = np.asarray(
        test_idx,
        dtype=np.int64,
    )

    X_train_raw = (
        X_raw[
            train_idx
        ]
    )

    X_test_raw = (
        X_raw[
            test_idx
        ]
    )

    y_train_raw = (
        y_raw[
            train_idx
        ]
    )

    y_test_raw = (
        y_raw[
            test_idx
        ]
    )

    feature_mean = np.mean(
        X_train_raw,
        axis=0,
    )

    feature_std = np.std(
        X_train_raw,
        axis=0,
        ddof=0,
    )

    zero_variance_mask = (
        feature_std == 0.0
    )

    kept_mask = (
        ~zero_variance_mask
    )

    if not np.any(
        kept_mask
    ):
        raise ValueError(
            f"{dataset.slug}: all "
            "training features have "
            "zero variance."
        )

    removed_features = [
        dataset.feature_names[index]
        for index
        in np.flatnonzero(
            zero_variance_mask
        )
    ]

    kept_feature_names = [
        dataset.feature_names[index]
        for index
        in np.flatnonzero(
            kept_mask
        )
    ]

    kept_mean = (
        feature_mean[
            kept_mask
        ]
    )

    kept_std = (
        feature_std[
            kept_mask
        ]
    )

    X_train = (
        X_train_raw[
            :,
            kept_mask,
        ]
        - kept_mean
    ) / kept_std

    X_test = (
        X_test_raw[
            :,
            kept_mask,
        ]
        - kept_mean
    ) / kept_std

    target_mean = float(
        np.mean(
            y_train_raw
        )
    )

    target_std = float(
        np.std(
            y_train_raw,
            ddof=0,
        )
    )

    if (
        not np.isfinite(
            target_std
        )
        or target_std == 0.0
    ):
        raise ValueError(
            f"{dataset.slug}: "
            "invalid target standard "
            f"deviation {target_std}."
        )

    y_train = (
        y_train_raw
        - target_mean
    ) / target_std

    y_test = (
        y_test_raw
        - target_mean
    ) / target_std

    for name, array in (
        (
            "X_train",
            X_train,
        ),
        (
            "X_test",
            X_test,
        ),
        (
            "y_train",
            y_train,
        ),
        (
            "y_test",
            y_test,
        ),
    ):
        assert_finite(
            f"{dataset.slug} "
            f"{name}",
            np.asarray(
                array
            ),
        )

    output_path = (
        PREPARED_ROOT
        / f"{dataset.slug}.npz"
    )

    np.savez_compressed(
        output_path,

        X_train=np.asarray(
            X_train,
            dtype=np.float64,
        ),

        y_train=np.asarray(
            y_train,
            dtype=np.float64,
        ),

        X_test=np.asarray(
            X_test,
            dtype=np.float64,
        ),

        y_test=np.asarray(
            y_test,
            dtype=np.float64,
        ),

        train_indices=(
            train_idx
        ),

        test_indices=(
            test_idx
        ),

        feature_names=np.asarray(
            kept_feature_names,
            dtype=np.str_,
        ),

        removed_feature_names=(
            np.asarray(
                removed_features,
                dtype=np.str_,
            )
        ),

        feature_mean=np.asarray(
            kept_mean,
            dtype=np.float64,
        ),

        feature_std=np.asarray(
            kept_std,
            dtype=np.float64,
        ),

        target_mean=np.asarray(
            target_mean,
            dtype=np.float64,
        ),

        target_std=np.asarray(
            target_std,
            dtype=np.float64,
        ),

        target_name=np.asarray(
            target_name,
            dtype=np.str_,
        ),
    )

    raw_content_hash = (
        hash_arrays_and_labels(
            (
                X_raw,
                y_raw,
            ),
            (
                *dataset.feature_names,
                target_name,
            ),
        )
    )

    prepared_content_hash = (
        hash_arrays_and_labels(
            (
                X_train,
                y_train,
                X_test,
                y_test,
                train_idx,
                test_idx,
            ),
            (
                *kept_feature_names,
                target_name,
            ),
        )
    )

    metadata = (
        REAL_DATASETS[
            dataset.slug
        ]
    )

    return {
        "kind": "real",

        "slug": (
            dataset.slug
        ),

        "display_name": (
            metadata[
                "display_name"
            ]
        ),

        "source": (
            metadata["source"]
        ),

        "source_id": (
            metadata[
                "source_id"
            ]
        ),

        "source_url": (
            metadata[
                "source_url"
            ]
        ),

        "doi": (
            metadata.get(
                "doi"
            )
        ),

        "loaded_from_local_raw_cache":
        dataset.loaded_from_cache,

        "raw_shape": [
            int(
                X_raw.shape[0]
            ),
            int(
                X_raw.shape[1]
            ),
        ],

        "raw_target_count": (
            int(
                dataset
                .Y
                .shape[1]
            )
        ),

        "returned_target_names": (
            list(
                dataset.target_names
            )
        ),

        "selected_target": (
            target_name
        ),

        "target_resolution": (
            target_resolution
        ),

        "split_seed": (
            SPLIT_SEED
        ),

        "train_rows": (
            int(
                X_train.shape[0]
            )
        ),

        "test_rows": (
            int(
                X_test.shape[0]
            )
        ),

        "original_feature_count": (
            int(
                X_raw.shape[1]
            )
        ),

        "prepared_feature_count": (
            int(
                X_train.shape[1]
            )
        ),

        "removed_zero_variance_features": (
            removed_features
        ),

        "kept_feature_names": (
            kept_feature_names
        ),

        "training_target_mean_raw": (
            target_mean
        ),

        "training_target_std_raw": (
            target_std
        ),

        "raw_content_sha256": (
            raw_content_hash
        ),

        "prepared_content_sha256": (
            prepared_content_hash
        ),

        "split_indices_sha256": (
            hash_arrays_and_labels(
                (
                    train_idx,
                    test_idx,
                )
            )
        ),

        "prepared_file": str(
            output_path
            .relative_to(ROOT)
        ),

        "prepared_file_sha256": (
            sha256_file(
                output_path
            )
        ),

        "raw_cache_file": str(
            raw_cache_path(
                dataset.slug
            )
            .relative_to(ROOT)
        ),

        "raw_cache_file_sha256": (
            sha256_file(
                raw_cache_path(
                    dataset.slug
                )
            )
        ),

        "source_metadata": (
            json_safe(
                dataset
                .source_metadata
            )
        ),
    }


# =============================================================================
# Synthetic generator
# =============================================================================

def toeplitz_covariance(
    d: int,
    rho: float,
) -> np.ndarray:
    indices = np.arange(
        d
    )

    distances = np.abs(
        indices[:, None]
        - indices[None, :]
    )

    covariance = (
        rho ** distances
    )

    return np.asarray(
        covariance,
        dtype=np.float64,
    )


def synthetic_substreams(
    seed: int,
) -> tuple[
    np.random.Generator,
    np.random.Generator,
    np.random.Generator,
]:
    seed_sequence = (
        np.random.SeedSequence(
            seed
        )
    )

    (
        beta_sequence,
        design_sequence,
        noise_sequence,
    ) = seed_sequence.spawn(
        3
    )

    return (
        np.random.default_rng(
            beta_sequence
        ),
        np.random.default_rng(
            design_sequence
        ),
        np.random.default_rng(
            noise_sequence
        ),
    )


def generate_synthetic_problem(
    case: dict[str, Any],
    seed: int,
) -> dict[str, Any]:

    case_name = str(
        case["name"]
    )

    rho = float(
        case["rho"]
    )

    scaled = bool(
        case["scaled"]
    )

    (
        beta_rng,
        design_rng,
        noise_rng,
    ) = synthetic_substreams(
        seed
    )

    beta = beta_rng.normal(
        size=SYNTHETIC_D
    )

    beta_norm = float(
        np.linalg.norm(
            beta
        )
    )

    if (
        not np.isfinite(
            beta_norm
        )
        or beta_norm == 0.0
    ):
        raise RuntimeError(
            f"{case_name}, "
            f"seed {seed}: "
            "invalid beta norm."
        )

    beta = (
        beta
        / beta_norm
    )

    covariance = (
        toeplitz_covariance(
            SYNTHETIC_D,
            rho,
        )
    )

    try:
        chol = (
            np.linalg.cholesky(
                covariance
            )
        )
    except np.linalg.LinAlgError as exc:
        raise RuntimeError(
            f"{case_name}: "
            "Toeplitz covariance "
            "is not positive definite."
        ) from exc

    Z = design_rng.normal(
        size=(
            SYNTHETIC_N,
            SYNTHETIC_D,
        )
    )

    X = (
        Z
        @ chol.T
    )

    if scaled:
        scales = np.logspace(
            SYNTHETIC_SCALE_MIN_EXP,
            SYNTHETIC_SCALE_MAX_EXP,
            SYNTHETIC_D,
        )
    else:
        scales = np.ones(
            SYNTHETIC_D,
            dtype=np.float64,
        )

    X = (
        X
        * scales
    )

    noise = noise_rng.normal(
        loc=0.0,
        scale=SYNTHETIC_NOISE_STD,
        size=SYNTHETIC_N,
    )

    y = (
        X @ beta
        + noise
    )

    X = np.asarray(
        X,
        dtype=np.float64,
    )

    y = np.asarray(
        y,
        dtype=np.float64,
    )

    beta = np.asarray(
        beta,
        dtype=np.float64,
    )

    scales = np.asarray(
        scales,
        dtype=np.float64,
    )

    assert_finite(
        f"{case_name} "
        f"seed={seed} X",
        X,
    )

    assert_finite(
        f"{case_name} "
        f"seed={seed} y",
        y,
    )

    output_path = (
        PREPARED_ROOT
        / (
            f"{case_name}"
            f"_seed{seed}.npz"
        )
    )

    np.savez_compressed(
        output_path,

        X=X,
        y=y,
        beta=beta,
        scales=scales,

        rho=np.asarray(
            rho,
            dtype=np.float64,
        ),

        seed=np.asarray(
            seed,
            dtype=np.int64,
        ),

        noise_std=np.asarray(
            SYNTHETIC_NOISE_STD,
            dtype=np.float64,
        ),
    )

    content_hash = (
        hash_arrays_and_labels(
            (
                X,
                y,
                beta,
                scales,
            ),
            (
                case_name,
                str(seed),
                f"{rho:.17g}",
                str(scaled),
            ),
        )
    )

    return {
        "kind": "synthetic",

        "case": case_name,

        "seed": seed,

        "n": SYNTHETIC_N,

        "d": SYNTHETIC_D,

        "rho": rho,

        "scaled": scaled,

        "scaling_rule": (
            "logspace(-1.5, 1.5, d)"
            if scaled
            else "all_ones"
        ),

        "beta_distribution": (
            SYNTHETIC_BETA_DISTRIBUTION
        ),

        "noise_distribution": (
            "normal"
        ),

        "noise_std": (
            SYNTHETIC_NOISE_STD
        ),

        "content_sha256": (
            content_hash
        ),

        "prepared_file": str(
            output_path
            .relative_to(ROOT)
        ),

        "prepared_file_sha256": (
            sha256_file(
                output_path
            )
        ),
    }


# =============================================================================
# Source schema validation
# =============================================================================

def validate_real_expected_schemas(
    records: list[
        dict[str, Any]
    ],
) -> None:
    by_slug = {
        record["slug"]:
        record
        for record
        in records
    }

    expected_slugs = set(
        EXPECTED_REAL_SCHEMAS
    )

    actual_slugs = set(
        by_slug
    )

    if (
        actual_slugs
        != expected_slugs
    ):
        raise AssertionError(
            "Real dataset set mismatch.\n"
            f"Expected: "
            f"{sorted(expected_slugs)}\n"
            f"Observed: "
            f"{sorted(actual_slugs)}"
        )

    for (
        slug,
        expected,
    ) in (
        EXPECTED_REAL_SCHEMAS
        .items()
    ):
        record = by_slug[
            slug
        ]

        actual_rows = (
            record[
                "raw_shape"
            ][0]
        )

        actual_features = (
            record[
                "raw_shape"
            ][1]
        )

        actual_targets = (
            record[
                "raw_target_count"
            ]
        )

        if (
            actual_rows
            != expected["rows"]
        ):
            raise AssertionError(
                f"{slug}: expected "
                f"{expected['rows']} rows; "
                f"found {actual_rows}."
            )

        if (
            actual_features
            != expected["features"]
        ):
            raise AssertionError(
                f"{slug}: expected "
                f"{expected['features']} "
                "predictor columns; "
                f"found {actual_features}."
            )

        if (
            actual_targets
            != expected["targets"]
        ):
            raise AssertionError(
                f"{slug}: expected "
                f"{expected['targets']} "
                "source target(s); "
                f"found {actual_targets}."
            )


# =============================================================================
# Manifest
# =============================================================================

def build_manifest(
    real_records: list[
        dict[str, Any]
    ],
    synthetic_records: list[
        dict[str, Any]
    ],
    command_line: list[str],
) -> dict[str, Any]:

    residential_schema_note = (
        "Official workbook inspection shows 109 Data-sheet columns: "
        "four project-date variables; eight project physical/financial "
        "variables V-1 through V-8; economic variables V-11 through "
        "V-29 repeated for time lags 1 through 5; and two outputs. "
        "V-9 is Actual sales prices (output) and V-10 is Actual "
        "construction costs (output). Both outputs are excluded from X, "
        "leaving 107 predictors. Lagged predictors are uniquely named "
        "V-11_lag1 through V-29_lag5. V-9 is the Stage 5 target."
    )

    return {
        "stage": 5,

        "purpose": (
            "confirmatory data acquisition "
            "and validation only"
        ),

        "generated_at_utc": (
            utc_now_string()
        ),

        "repository_root": (
            str(ROOT)
        ),

        "command_line": (
            command_line
        ),

        "protocol": {
            "split_seed": (
                SPLIT_SEED
            ),

            "train_fraction": (
                TRAIN_FRACTION
            ),

            "test_fraction": (
                TEST_FRACTION
            ),

            "real_preprocessing": {
                "feature_center": (
                    "training_mean"
                ),

                "feature_scale": (
                    "training_population_"
                    "std_ddof0"
                ),

                "zero_variance_policy": (
                    "remove_and_record"
                ),

                "target_center": (
                    "training_mean"
                ),

                "target_scale": (
                    "training_population_"
                    "std_ddof0"
                ),

                "test_statistics_source": (
                    "training_only"
                ),
            },

            "residential_schema_clarification":
            residential_schema_note,

            "synthetic": {
                "n": (
                    SYNTHETIC_N
                ),

                "d": (
                    SYNTHETIC_D
                ),

                "seeds": list(
                    SYNTHETIC_SEEDS
                ),

                "covariance": (
                    "Toeplitz: "
                    "rho^|i-j|"
                ),

                "rho_values": [
                    0.60,
                    0.90,
                ],

                "scaled_feature_rule": (
                    "logspace("
                    "-1.5, 1.5, d)"
                ),

                "beta_distribution": (
                    SYNTHETIC_BETA_DISTRIBUTION
                ),

                "noise_distribution": (
                    "normal"
                ),

                "noise_std": (
                    SYNTHETIC_NOISE_STD
                ),
            },
        },

        "environment": {
            "python": (
                sys.version
            ),

            "python_executable": (
                sys.executable
            ),

            "platform": (
                platform.platform()
            ),

            "numpy": (
                np.__version__
            ),

            "pandas": (
                pd.__version__
            ),

            "scikit_learn": (
                sklearn.__version__
            ),

            "ucimlrepo": (
                package_version(
                    "ucimlrepo"
                )
            ),

            "openpyxl": (
                package_version(
                    "openpyxl"
                )
            ),
        },

        "real_datasets": (
            real_records
        ),

        "synthetic_datasets": (
            synthetic_records
        ),

        "validation": {
            "finite_values_checked":
            True,

            "real_expected_schemas_checked":
            True,

            "residential_sale_price_verified_from_official_workbook_description":
            True,

            "residential_other_output_excluded_from_predictors":
            True,

            "residential_predictor_names_unique":
            True,

            "residential_lag_structure_validated":
            True,

            "optimization_methods_run":
            False,
        },
    }


# =============================================================================
# CLI
# =============================================================================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Prepare and validate Stage 5 "
            "confirmatory datasets without "
            "running optimization."
        )
    )

    parser.add_argument(
        "--refresh",
        action="store_true",
        help=(
            "Refetch external datasets "
            "instead of using Stage 5 "
            "raw/download caches."
        ),
    )

    return parser.parse_args()


# =============================================================================
# Main
# =============================================================================

def main() -> None:
    args = parse_args()

    ensure_directories()

    print(
        "=" * 72
    )
    print(
        "STAGE 5 DATA PREPARATION PREFLIGHT"
    )
    print(
        "=" * 72
    )

    print(
        "\nThis script will not run "
        "any optimization method."
    )

    print(
        "\nFrozen preparation parameters:"
    )

    print(
        f"  split seed: "
        f"{SPLIT_SEED}"
    )

    print(
        f"  real train/test: "
        f"{TRAIN_FRACTION:.0%}/"
        f"{TEST_FRACTION:.0%}"
    )

    print(
        f"  synthetic shape: "
        f"{SYNTHETIC_N} x "
        f"{SYNTHETIC_D}"
    )

    print(
        f"  synthetic seeds: "
        f"{SYNTHETIC_SEEDS[0]}"
        f".."
        f"{SYNTHETIC_SEEDS[-1]}"
    )

    print(
        f"  synthetic noise std: "
        f"{SYNTHETIC_NOISE_STD}"
    )

    print(
        f"  synthetic beta: "
        f"{SYNTHETIC_BETA_DISTRIBUTION}"
    )

    real_records: list[
        dict[str, Any]
    ] = []

    # -------------------------------------------------------------------------
    # Real 1: California Housing
    # -------------------------------------------------------------------------

    print(
        "\n[REAL 1/4] "
        "California Housing"
    )

    california = load_california(
        refresh=args.refresh
    )

    (
        y,
        target_name,
        resolution,
    ) = resolve_single_target(
        california
    )

    record = prepare_real_dataset(
        california,
        y,
        target_name,
        resolution,
    )

    real_records.append(
        record
    )

    print(
        "  prepared "
        f"{record['raw_shape']} "
        f"-> train="
        f"{record['train_rows']}, "
        f"test="
        f"{record['test_rows']}, "
        f"d="
        f"{record['prepared_feature_count']}"
    )

    # -------------------------------------------------------------------------
    # Real 2: Concrete
    # -------------------------------------------------------------------------

    print(
        "\n[REAL 2/4] "
        "Concrete Compressive Strength"
    )

    concrete = load_uci_dataset(
        slug=(
            "concrete_compressive_strength"
        ),
        uci_id=165,
        refresh=args.refresh,
    )

    (
        y,
        target_name,
        resolution,
    ) = resolve_single_target(
        concrete
    )

    record = prepare_real_dataset(
        concrete,
        y,
        target_name,
        resolution,
    )

    real_records.append(
        record
    )

    print(
        "  prepared "
        f"{record['raw_shape']} "
        f"-> train="
        f"{record['train_rows']}, "
        f"test="
        f"{record['test_rows']}, "
        f"d="
        f"{record['prepared_feature_count']}"
    )

    # -------------------------------------------------------------------------
    # Real 3: Airfoil
    # -------------------------------------------------------------------------

    print(
        "\n[REAL 3/4] "
        "Airfoil Self-Noise"
    )

    airfoil = load_uci_dataset(
        slug=(
            "airfoil_self_noise"
        ),
        uci_id=291,
        refresh=args.refresh,
    )

    (
        y,
        target_name,
        resolution,
    ) = resolve_single_target(
        airfoil
    )

    record = prepare_real_dataset(
        airfoil,
        y,
        target_name,
        resolution,
    )

    real_records.append(
        record
    )

    print(
        "  prepared "
        f"{record['raw_shape']} "
        f"-> train="
        f"{record['train_rows']}, "
        f"test="
        f"{record['test_rows']}, "
        f"d="
        f"{record['prepared_feature_count']}"
    )

    # -------------------------------------------------------------------------
    # Real 4: Residential Building
    # -------------------------------------------------------------------------

    print(
        "\n[REAL 4/4] "
        "Residential Building"
    )

    residential = (
        load_residential_direct(
            refresh=args.refresh
        )
    )

    print(
        "  returned output columns: "
        f"{residential.target_names}"
    )

    for target in (
        residential.target_names
    ):
        description = (
            residential
            .target_descriptions
            .get(
                target,
                "",
            )
        )

        print(
            f"    {target!r}: "
            f"{description!r}"
        )

    (
        y,
        target_name,
        resolution,
    ) = (
        resolve_residential_sale_target(
            residential
        )
    )

    print(
        "  selected Stage 5 "
        "sale-price target: "
        f"{target_name!r}"
    )

    print(
        "  excluded from predictors: "
        "'V-9' and 'V-10'"
    )

    record = prepare_real_dataset(
        residential,
        y,
        target_name,
        resolution,
    )

    real_records.append(
        record
    )

    print(
        "  prepared "
        f"{record['raw_shape']} "
        f"-> train="
        f"{record['train_rows']}, "
        f"test="
        f"{record['test_rows']}, "
        f"d="
        f"{record['prepared_feature_count']}"
    )

    # -------------------------------------------------------------------------
    # Validate all real source schemas
    # -------------------------------------------------------------------------

    validate_real_expected_schemas(
        real_records
    )

    print(
        "\nReal-data schema validation: PASSED"
    )

    # -------------------------------------------------------------------------
    # Held-out synthetic problems
    # -------------------------------------------------------------------------

    print(
        "\n[HELD-OUT SYNTHETIC]"
    )

    print(
        "Generating "
        f"{len(SYNTHETIC_CASES)} "
        "cases x "
        f"{len(SYNTHETIC_SEEDS)} "
        "seeds"
    )

    synthetic_records: list[
        dict[str, Any]
    ] = []

    for case in (
        SYNTHETIC_CASES
    ):
        print(
            "\n  "
            f"{case['name']} "
            f"(rho={case['rho']}, "
            f"scaled="
            f"{case['scaled']})"
        )

        for seed in (
            SYNTHETIC_SEEDS
        ):
            synthetic_record = (
                generate_synthetic_problem(
                    case,
                    seed,
                )
            )

            synthetic_records.append(
                synthetic_record
            )

            print(
                f"    seed {seed}: "
                f"{synthetic_record['content_sha256'][:12]}"
                "..."
            )

    expected_synthetic_count = (
        len(
            SYNTHETIC_CASES
        )
        * len(
            SYNTHETIC_SEEDS
        )
    )

    if (
        len(
            synthetic_records
        )
        != expected_synthetic_count
    ):
        raise AssertionError(
            "Synthetic output count "
            "mismatch: "
            f"{len(synthetic_records)} "
            "!= "
            f"{expected_synthetic_count}"
        )

    print(
        "\nSynthetic artifact "
        "count validation: PASSED "
        f"({expected_synthetic_count})"
    )

    # -------------------------------------------------------------------------
    # Final manifest
    # -------------------------------------------------------------------------

    manifest = build_manifest(
        real_records=(
            real_records
        ),
        synthetic_records=(
            synthetic_records
        ),
        command_line=(
            sys.argv
        ),
    )

    temporary_manifest = (
        MANIFEST_PATH
        .with_suffix(
            ".json.tmp"
        )
    )

    temporary_manifest.write_text(
        json.dumps(
            manifest,
            indent=2,
            sort_keys=False,
        )
        + "\n",
        encoding="utf-8",
    )

    # Parse from disk before replacing
    # the canonical manifest.
    json.loads(
        temporary_manifest.read_text(
            encoding="utf-8"
        )
    )

    temporary_manifest.replace(
        MANIFEST_PATH
    )

    # Final independent parse.
    json.loads(
        MANIFEST_PATH.read_text(
            encoding="utf-8"
        )
    )

    print(
        "\n"
        + "=" * 72
    )

    print(
        "STAGE 5 DATA PREPARATION: PASSED"
    )

    print(
        "=" * 72
    )

    print(
        "\nReal datasets prepared: "
        f"{len(real_records)}"
    )

    print(
        "Synthetic problems prepared: "
        f"{len(synthetic_records)}"
    )

    print(
        "\nResidential Building:"
    )

    print(
        "  source rows: 372"
    )

    print(
        "  predictors before "
        "zero-variance filtering: 107"
    )

    print(
        "  predictor identifiers: "
        "unique, lag-aware"
    )

    print(
        "  lagged variables: "
        "V-11_lag1 ... V-29_lag5"
    )

    print(
        "  target: V-9 "
        "(Actual sales prices)"
    )

    print(
        "  excluded second output: "
        "V-10 "
        "(Actual construction costs)"
    )

    print(
        "\nManifest:"
    )

    print(
        f"  {MANIFEST_PATH}"
    )

    print(
        "\nPrepared data:"
    )

    print(
        f"  {PREPARED_ROOT}"
    )

    print(
        "\nLocal raw cache:"
    )

    print(
        f"  {RAW_CACHE_ROOT}"
    )

    print(
        "\nNO OPTIMIZATION METHODS WERE RUN."
    )

    print(
        "\nReview dataset_manifest.json "
        "before implementing or running "
        "the Stage 5 benchmark."
    )


if __name__ == "__main__":
    main()
